"""Shared builders for test documents."""

import io

from docx import Document as NewDocx
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from sharesentry.parsers import parse_document

EMAIL = "jordan.rivera@example-client.test"
CARD = "4111 1111 1111 1111"
TOKEN = "sk-FAKE0000000000000000000000EXAMPLE"


def parse(name: str, data: bytes):
    return parse_document(name, data, max_bytes=len(data) + 1)


def make_docx(paragraphs, table_cell=None) -> bytes:
    document = NewDocx()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_cell:
        document.add_table(rows=1, cols=1).cell(0, 0).text = table_cell
    document.sections[0].header.paragraphs[0].text = "Header text"
    document.core_properties.author = "Secret Author"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_pdf(pages) -> bytes:
    """Each entry is page text; an empty string yields a blank (image-only style) page."""
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=A4)
    for text in pages:
        for i, line in enumerate(text.splitlines()):
            pdf.drawString(50, 800 - 14 * i, line)
        pdf.showPage()
    pdf.save()
    return out.getvalue()
