import csv
import io
import unittest

from docx import Document as OpenDocx

from sharesentry.audit import run_audit, verify_redaction
from sharesentry.models import Context
from sharesentry.parsers import extension, parse_document
from sharesentry.redaction import redact_document
from sharesentry.skills import SKILLS_BY_NAME
from tests.helpers import CARD, EMAIL, TOKEN, make_docx, make_pdf, parse

SKILL = SKILLS_BY_NAME["Client-Ready Review"]


def audit_and_redact(doc, **kwargs):
    audit = run_audit(doc, SKILL, Context.EXTERNAL_CLIENT)
    result = redact_document(doc, audit.findings, {f.finding_id for f in audit.findings}, **kwargs)
    return audit, result


class ParserTests(unittest.TestCase):
    def test_text_markdown_and_encodings(self):
        self.assertEqual(parse("a.md", b"# Title").fmt, "md")
        latin = parse("a.txt", "caf\xe9 hello".encode("latin-1"))
        self.assertEqual(latin.text, "caf\xe9 hello")
        self.assertTrue(any("Latin-1" in n for n in latin.scope_notes))
        self.assertEqual(parse("a.txt", "\ufeffbom".encode("utf-8")).text, "bom")

    def test_unsupported_empty_and_oversize_are_incomplete(self):
        self.assertIn("Unsupported", parse("a.exe", b"MZ").incomplete_reasons[0])
        self.assertIn("empty", parse("a.txt", b"  \n").incomplete_reasons[0])
        big = parse_document("a.txt", b"x" * 20, max_bytes=10)
        self.assertIn("limit", big.incomplete_reasons[0])
        self.assertFalse(big.inspectable)
        self.assertEqual(extension("Report.PDF"), "pdf")

    def test_csv_cells_become_segments_and_delimiter_is_detected(self):
        doc = parse("a.csv", b"name;email\nJo;jo@example.test\n")
        self.assertEqual(doc.meta["delimiter"], ";")
        self.assertEqual([s.ref for s in doc.segments][-1], ("csv", 1, 1))

    def test_docx_covers_body_tables_and_headers(self):
        doc = parse("a.docx", make_docx(["body one", f"mail {EMAIL}"], table_cell="table cell"))
        texts = [s.text for s in doc.segments]
        for expected in ("body one", "table cell", "Header text"):
            self.assertIn(expected, texts)
        self.assertTrue(any("comments" in n for n in doc.scope_notes))

    def test_fake_or_corrupt_files_are_incomplete_not_crashes(self):
        self.assertTrue(parse("a.docx", b"not a zip").incomplete_reasons)
        self.assertTrue(parse("a.docx", b"PK\x03\x04garbage").incomplete_reasons)
        self.assertTrue(parse("a.pdf", b"nope").incomplete_reasons)
        self.assertTrue(parse("a.pdf", b"%PDF-1.4 garbage").incomplete_reasons)

    def test_pdf_text_and_image_only_pages(self):
        good = parse("a.pdf", make_pdf([f"hello {EMAIL}", "second page"]))
        self.assertEqual(len(good.segments), 2)
        self.assertFalse(good.incomplete_reasons)
        partial = parse("a.pdf", make_pdf(["has text", ""]))
        self.assertIn("Page(s) 2", partial.incomplete_reasons[0])
        scanned = parse("a.pdf", make_pdf(["", ""]))
        self.assertFalse(scanned.inspectable)
        self.assertTrue(scanned.incomplete_reasons)


class RedactionTests(unittest.TestCase):
    def test_text_redaction_replaces_exact_spans_and_verifies_clean(self):
        doc = parse("a.txt", f"mail {EMAIL}\ncard {CARD}\nkey {TOKEN}\nkeep this".encode())
        audit, result = audit_and_redact(doc)
        out = result.data.decode()
        self.assertEqual(out, "mail [REDACTED:EMAIL]\ncard [REDACTED:CREDIT_CARD]\nkey [REDACTED:API_KEY]\nkeep this")
        self.assertTrue(verify_redaction(result, audit, SKILL, []).clean)

    def test_partial_selection_leaves_unselected_findings_and_reports_them(self):
        doc = parse("a.txt", f"{EMAIL} and {TOKEN}".encode())
        audit = run_audit(doc, SKILL, Context.EXTERNAL_CLIENT)
        token = next(f for f in audit.findings if f.category.value == "api_key")
        result = redact_document(doc, audit.findings, {token.finding_id})
        verification = verify_redaction(result, audit, SKILL, [])
        self.assertEqual([f.category.value for f in verification.remaining], ["email"])
        self.assertEqual(verification.redacted_count, 1)

    def test_csv_preserves_shape_and_can_drop_columns(self):
        data = f"name,email,note\nJo,{EMAIL},hi\nSam,,ok\n".encode()
        doc = parse("a.csv", data)
        _, result = audit_and_redact(doc)
        rows = list(csv.reader(io.StringIO(result.data.decode())))
        self.assertEqual(rows, [["name", "email", "note"], ["Jo", "[REDACTED:EMAIL]", "hi"], ["Sam", "", "ok"]])
        _, dropped = audit_and_redact(doc, drop_columns=[1])
        self.assertEqual(list(csv.reader(io.StringIO(dropped.data.decode())))[1], ["Jo", "hi"])

    def test_docx_keeps_structure_and_removes_secrets(self):
        source = make_docx(["Intro", f"Contact {EMAIL} now", "Outro"], table_cell=f"card {CARD}")
        doc = parse("a.docx", source)
        audit, result = audit_and_redact(doc)
        out = OpenDocx(io.BytesIO(result.data))
        paragraphs = [p.text for p in out.paragraphs]
        self.assertEqual(paragraphs[:3], ["Intro", "Contact [REDACTED:EMAIL] now", "Outro"])
        self.assertEqual(out.tables[0].cell(0, 0).text, "card [REDACTED:CREDIT_CARD]")
        self.assertEqual(out.core_properties.author, "")
        self.assertTrue(verify_redaction(result, audit, SKILL, []).clean)

    def test_pdf_redaction_output_is_a_valid_pdf_without_findings(self):
        doc = parse("a.pdf", make_pdf([f"hello {EMAIL}", f"card {CARD}"]))
        audit, result = audit_and_redact(doc)
        self.assertTrue(result.data.startswith(b"%PDF"))
        reread = parse(result.name, result.data)
        self.assertNotIn("example-client", reread.text)
        self.assertIn("REDACTED:EMAIL", reread.text)
        self.assertTrue(verify_redaction(result, audit, SKILL, []).clean)

    def test_multi_segment_finding_is_removed_from_every_segment(self):
        key = ["-----BEGIN RSA PRIVATE KEY-----", "MIIFAKEFAKE", "-----END RSA PRIVATE KEY-----", "after"]
        doc = parse("a.docx", make_docx(["before", *key]))
        audit, result = audit_and_redact(doc)
        out = [p.text for p in OpenDocx(io.BytesIO(result.data)).paragraphs]
        self.assertEqual(out[0], "before")
        self.assertIn("[REDACTED:PRIVATE_KEY]", out[1])
        self.assertNotIn("MIIFAKEFAKE", "".join(out))
        self.assertEqual(out[-1], "after")


if __name__ == "__main__":
    unittest.main()
